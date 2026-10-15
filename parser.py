"""Safe, traceable field extraction. External LLM is always opt-in."""
import io
import json
import os
import re
import urllib.request
import zipfile
from pathlib import Path
from seumteo_layout import parse_seumteo

ALIAS = {
 'permit_type': ['신청 구분', '허가 신고 구분', '신청종류', '민원구분'],
 'work_type': ['건축 구분', '건축구분', '건축종별'],
 'site_address': ['대지 위치', '대지위치', '대지의 위치', '대지 소재지', '소재지'],
 'site_area_m2': ['대지면적', '대지 면적', '대지면적(㎡)'],
 'building_area_m2': ['건축면적', '건축 면적'],
 'total_floor_area_m2': ['연면적', '연 면적'],
 'building_coverage_percent': ['건폐율'],
 'floor_area_ratio_percent': ['용적률'],
 'primary_use': ['주용도', '주 용도', '건축물용도', '건축물 용도', '주된 용도', '주용도(건축물)'],
 'structure': ['주구조', '구조'],
 'above_ground_floors': ['지상층수', '지상 층수', '지상층'],
 'underground_floors': ['지하층수', '지하 층수', '지하층'],
 'building_count': ['건축물 동수', '동수', '건축물 수'],
 'parking_spaces': ['주차계획', '주차대수', '주차 대수'],
 'zoning': ['용도지역', '용도 지역'],
 'notes': ['특이사항', '특기사항']
}


def _decode(data):
    for codec in ('utf-8-sig', 'cp949', 'euc-kr'):
        try:
            return data.decode(codec)
        except UnicodeDecodeError:
            pass
    raise ValueError('문자 인코딩을 확인할 수 없습니다. UTF-8 또는 CP949 파일을 사용하세요.')


def _zip_safety(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            files = z.infolist()
            if len(files) > 400 or sum(x.file_size for x in files) > 30 * 1024 * 1024:
                raise ValueError('압축된 문서의 크기 또는 항목 수가 제한을 초과했습니다.')
    except zipfile.BadZipFile as e:
        raise ValueError('정상적인 문서 압축파일이 아닙니다.') from e


def extract_lines(filename, data):
    """Return [{'n':int, 'text':str, 'where':str}]. A source is not invented."""
    ext = Path(filename).suffix.lower()
    out = []

    def append_text(text, location):
        for line in text.splitlines():
            if line.strip():
                out.append({'n': len(out) + 1, 'text': line.strip()[:1000], 'where': location})
                if len(out) >= 15000:
                    raise ValueError('문서가 너무 깁니다. 15,000행 이내 자료를 사용하세요.')

    if ext in ('.txt', '.csv'):
        append_text(_decode(data), '텍스트')
    elif ext == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ValueError('PDF 분석 모듈이 없습니다. install_parsers_windows.bat을 실행하세요.') from e
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError('암호화 PDF는 지원하지 않습니다.')
        if len(reader.pages) > 80:
            raise ValueError('PDF는 최대 80쪽까지 처리할 수 있습니다.')
        for i, page in enumerate(reader.pages, start=1):
            append_text(page.extract_text() or '', f'{i}쪽')
            if i == 1:
                # Seumteo ticks can be drawn as PDF vectors and are invisible
                # to ordinary text extraction. Never guess if ambiguous.
                try:
                    from pdf_vectors import checked_work_type
                    chosen = checked_work_type(page)
                except (ValueError, TypeError, IndexError):
                    chosen = None
                if chosen:
                    # A clearly labelled graphic-derived inspection record,
                    # not a literal quote from the document's text layer.
                    out.append({
                        'n': len(out) + 1,
                        'text': '건축구분: ' + chosen['value'],
                        'where': '1쪽 · PDF 벡터 체크표시',
                        'evidence_note': chosen['evidence'],
                        'kind': 'graphic_inspection'
                    })
    elif ext == '.docx':
        _zip_safety(data)
        try:
            from docx import Document
        except ImportError as e:
            raise ValueError('DOCX 분석 모듈이 없습니다. install_parsers_windows.bat을 실행하세요.') from e
        doc = Document(io.BytesIO(data))
        for i, p in enumerate(doc.paragraphs, 1):
            append_text(p.text, f'문단 {i}')
        for it, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                cells = [cell.text.strip() for cell in row.cells]
                append_text(': '.join(x for x in cells if x), f'표 {it}·행 {r}')
    elif ext == '.xls':
        from xls_legacy import read_xls_cells
        from collections import defaultdict
        cell_rows = defaultdict(list)
        for sheet, row_number, col, value in read_xls_cells(data):
            cell_rows[(sheet, row_number)].append({'col':col,'text':value})
            if len(cell_rows) > 15000:
                raise ValueError('XLS 분석은 15,000행 이내로 제한합니다.')
        for (sheet, row_number), cells in cell_rows.items():
            cells = sorted(cells, key=lambda c: c['col'])
            text = ' | '.join(c['text'] for c in cells)
            append_text(text, f'{sheet} {row_number}행')
            out[-1]['cells'] = cells
        from xls_markers import checked_work_type
        chosen = checked_work_type(data)
        if chosen:
            # The visible check is a picture object and was absent from the
            # original extracted text. Mark the derived row transparently.
            out.append({
                'n': len(out) + 1,
                'text': '건축구분: ' + chosen['value'],
                'where': 'Page 1 · XLS 그림 체크표시',
                'evidence_note': chosen['evidence'],
                'kind': 'graphic_inspection'
            })
    elif ext == '.xlsx':
        _zip_safety(data)
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise ValueError('XLSX 분석 모듈이 없습니다. install_parsers_windows.bat을 실행하세요.') from e
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for sheet in wb:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [str(v).strip() for v in row if v is not None]
                if cells:
                    append_text(': '.join(cells), f'{sheet.title} {row_number}행')
                if row_number > 10000:
                    raise ValueError('시트는 10,000행 이내로 제한합니다.')
        wb.close()
    else:
        raise ValueError('지원 형식은 TXT, CSV, PDF(텍스트형), DOCX, XLS, XLSX입니다. HWP/HWPX는 현재 미지원입니다.')
    if not out:
        raise ValueError('읽을 수 있는 글자가 없습니다. 스캔 PDF 등은 현재 자동 추출할 수 없습니다.')
    return out


def _clean(key, value):
    val = str(value or '').strip().strip('·•').strip()
    if key in ('site_area_m2', 'building_area_m2', 'total_floor_area_m2',
               'building_coverage_percent', 'floor_area_ratio_percent'):
        m = re.search(r'([\d,]+(?:\.\d+)?)', val)
        return m.group(1).replace(',', '') if m else ''
    if key in ('above_ground_floors','underground_floors','building_count','parking_spaces'):
        m = re.search(r'(\d+)', val.replace(',', ''))
        return m.group(1) if m else ''
    if key == 'permit_type':
        if val.strip() in ('건축신고', '신고'): return '건축신고'
        if val.strip() in ('건축허가', '허가'): return '건축허가'
        return ''
    return val[:1500]


def local_parse(lines, allowed_keys):
    """Extract only labeled entries, preserving the precise source line.

    Supports both ``대지면적: 350㎡`` and ``대지면적(㎡) 350.00``.
    Does not infer a permit type from a heading, or guess from other fields.
    """
    values, sources = {}, {}
    aliases = sorted(((alias, key) for key, lst in ALIAS.items()
                      for alias in lst if key in allowed_keys),
                     key=lambda x: -len(x[0]))
    number_keys = {'site_area_m2', 'building_area_m2', 'total_floor_area_m2',
                   'building_coverage_percent', 'floor_area_ratio_percent',
                   'above_ground_floors', 'underground_floors', 'building_count', 'parking_spaces'}
    for line in lines:
        text = line['text'].strip()
        # Splitting table cells at tabs or pipes keeps one field per fragment.
        fragments = re.split(r'\s*\|\s*|\t+', text)
        for fragment in fragments:
            s = fragment.strip()
            for alias, key in aliases:
                if key in values:
                    continue
                # Ignore explanatory prose, require label at the start of a fragment.
                pattern = (r'^\s*(?:[\[【(（]?\d+[)）.\]]?\s*)?'
                           + re.escape(alias) + r'\s*(?:\((?:㎡|m2|%|대|동)\))?'
                           + r'(?:\s*[:：=]\s*|\s+)(.+)$')
                match = re.match(pattern, s, flags=re.IGNORECASE)
                if not match:
                    continue
                candidate = match.group(1).strip()
                # Numbers must begin with the actual quantity; never pick the
                # first digit of another label or a date later in the line.
                if key in number_keys and not re.match(r'^[+]?\d[\d,]*(?:\.\d+)?\s*(?:㎡|m²|m2|%|층|동|대)?(?:\s|$)', candidate):
                    continue
                value = _clean(key, candidate)
                if value:
                    values[key] = value
                    sources[key] = {'line': line['n'], 'where': line['where'],
                                    'evidence': s, 'method': '규칙 기반', 'edited_by_user': False}
                break
    return values, sources


def ai_available():
    return (os.environ.get('AI_MODE', '').lower() == 'on'
            and bool(os.environ.get('AI_API_KEY'))
            and bool(os.environ.get('AI_MODEL')))


def ai_parse(lines, allowed_keys):
    """OpenAI-compatible chat API, without provider lock-in. Never run without consent."""
    if not ai_available():
        raise ValueError('AI 연결이 설정되지 않았습니다.')
    base = os.environ.get('AI_BASE_URL', 'https://api.openai.com/v1').rstrip('/')
    if not base.startswith('https://') and not base.startswith('http://127.0.0.1:') and not base.startswith('http://localhost:'):
        raise ValueError('AI_BASE_URL은 HTTPS 또는 로컬 주소만 허용합니다.')
    # Limit external exposure; model receives only extracted text, not original file bytes.
    numbered = '\n'.join(f"[{x['n']}] {x['text']}" for x in lines)
    numbered = numbered[:24000]
    prompt = (
        '당신은 건축허가·신고 신청서 데이터 추출기입니다. 행정적 적합성 판단을 하지 마세요. '
        '입력에 명시되지 않은 값은 추정하지 마세요. 출력은 JSON 객체만 반환하세요. '
        '필드 키는 허용키에 있는 것만 사용하세요. '\
        '형식: {"fields":{"key":"value"},"source_lines":{"key":줄번호}}. '
        '면적·비율은 단위를 제외한 숫자, 층수·동수·대수는 정수 문자열입니다.\n'
        f'허용키: {", ".join(sorted(allowed_keys))}\n입력:\n{numbered}'
    )
    payload = {'model': os.environ['AI_MODEL'], 'temperature': 0,
               'messages': [{'role': 'user', 'content': prompt}]}
    req = urllib.request.Request(
        base + '/chat/completions',
        data=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
        headers={'Content-Type': 'application/json', 'Authorization': 'Bearer ' + os.environ['AI_API_KEY']},
        method='POST'
    )
    with urllib.request.urlopen(req, timeout=25) as resp:
        obj = json.loads(resp.read(1024*1024).decode('utf-8'))
    raw = obj['choices'][0]['message']['content']
    raw = re.sub(r'^```(?:json)?\s*|\s*```$', '', raw.strip(), flags=re.IGNORECASE)
    decoded = json.loads(raw)
    allowed_sources = {x['n']: x for x in lines}
    fields, sources = {}, {}
    if not isinstance(decoded.get('fields'), dict):
        raise ValueError('AI 응답에 fields 객체가 없습니다.')
    for key, val in decoded['fields'].items():
        if key not in allowed_keys or not isinstance(val, (str, int, float)):
            continue
        clean = _clean(key, val)
        if not clean:
            continue
        line_num = decoded.get('source_lines', {}).get(key)
        try:
            cited = allowed_sources.get(int(line_num))
        except (TypeError, ValueError):
            cited = None
        # No source claim unless mapped to a real line.
        fields[key] = clean
        sources[key] = {'line': cited['n'] if cited else None,
                        'where': cited['where'] if cited else '',
                        'evidence': cited['text'] if cited else '원문 줄을 특정하지 못함',
                        'method': '외부 AI 초안', 'edited_by_user': False}
    return fields, sources


def parse_application(filename, data, allowed_keys, consent=False):
    lines = extract_lines(filename, data)
    values, sources = local_parse(lines, allowed_keys)
    form_values, form_sources = parse_seumteo(lines, allowed_keys)
    # The overall form summary overrides later building/floor details.
    values.update(form_values)
    sources.update(form_sources)
    method, warning = '규칙 기반', ''
    # AI runs separately through /api/ai-review and never silently overwrites values.
    return lines, values, sources, method, warning
